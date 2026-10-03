#!/usr/bin/env python3
"""Gera dados/Recife_Resiliente_Checklist_Subcriterios.xlsx.

Uma linha por subcritério: como medir, escala de notas de 1 a 5, origem da
faixa, recomendação, fonte e checklist de confirmação dos dados. Editar a
lista SUBCRITERIOS abaixo e rodar de novo para atualizar a planilha.
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SAIDA = Path(__file__).resolve().parent.parent / "dados" / "Recife_Resiliente_Checklist_Subcriterios.xlsx"
FONTE = "Arial"
PONTOS = SAIDA.parent / "processados" / "pontos_rpa6.csv"
Q = ("Quintil 1 (20% menores)", "Quintil 2", "Quintil 3", "Quintil 4", "Quintil 5 (20% maiores)")

# (critério, código, subcritério, como medir, tipo de nota, (nota 1..5), origem da faixa,
#  recomendação, fonte, observação, busca prévia)
SUBCRITERIOS = [
    ("Probabilidade", "P1", "Frequência e recorrência histórica",
     "Média de eventos por ano nas 3 últimas estações chuvosas", "Faixa fixa",
     ("< 1", "1–2", "3–4", "5–6", "≥ 7"), "Planilha original (provisória)",
     "Manter — peso alto",
     "Ficha preenchida com a EMLURB (agora); ocorrências do COP e da Defesa Civil (confirmação)",
     "Pedido ao COP via tutora/CGM. 156 só para conferir. Calibrar faixas com a distribuição real. "
     "'Recorrência histórica' (documento da equipe) fica aqui se for o mesmo que frequência (pendência 16).",
     "Não verificado: depende de pedido aos órgãos."),
    ("Probabilidade", "P2", "Severidade",
     "Altura máxima da água e tempo para baixar; nota = média das duas notas, arredondada para cima",
     "Faixa fixa",
     ("Altura < 10 cm / baixa em < 1 h", "10–30 cm (tornozelo) / 1–3 h",
      "30–50 cm (joelho) / 3–6 h", "50–100 cm (cintura) / 6–12 h", "> 100 cm / > 12 h"),
     "Planilha original (provisória); regra da média: proposta Claude",
     "Manter — peso alto", "Ficha EMLURB; escuta com Defesa Civil e moradores",
     "Extensão da mancha fica registrada como informação, sem nota própria.",
     "Não verificado: dado de campo."),
    ("Probabilidade", "P3", "Influência da maré",
     "O escoamento trava com a maré alta?", "Classe",
     ("Não influencia", "—", "Agrava ocasionalmente", "—",
      "Alaga dependendo da maré / escoamento bloqueado"),
     "Planilha original", "Manter", "Ficha EMLURB; APAC; tábua de marés",
     "Notas 2 e 4 ficam para casos intermediários, se a EMLURB indicar.",
     "APAC confirmada na busca; tábua de marés não verificada."),
    ("Probabilidade", "P4", "Proximidade de rio ou canal",
     "Distância do ponto ao curso d'água mais próximo (m); quanto mais perto, maior a nota",
     "Quintil (invertido)",
     ("Quintil mais distante", "Quintil 2", "Quintil 3", "Quintil 4", "Quintil mais próximo"),
     "Regra de quintis da planilha original",
     "Manter — peso menor (suscetibilidade)",
     "Shapefile próprio da rede hidrográfica; ESIG – recursos hídricos",
     "Reprojetar para EPSG:31985. Conferir se os canais estão no arquivo.",
     "ESIG com recursos hídricos confirmado na busca (não aberto)."),
    ("Probabilidade", "P5", "Impermeabilização",
     "% do círculo de 300 m com vegetação ou solo permeável; quanto menos verde, maior a nota",
     "Faixa fixa", ("> 40%", "30–40%", "20–30%", "10–20%", "< 10%"),
     "Planilha original (provisória)", "Manter — peso menor (suscetibilidade)",
     "MapBiomas (30 m); alternativa: edificações do ESIG",
     "Se quase todos os pontos ficarem em < 10%, trocar por quintis.", "Não verificado."),
    ("Probabilidade", "P6", "Relevo (baixio)",
     "Cota do ponto menos a cota média do círculo (m); quanto mais baixo, maior a nota",
     "Quintil (invertido)",
     ("Quintil mais alto", "Quintil 2", "Quintil 3", "Quintil 4", "Quintil mais baixo"),
     "Proposta Claude (quintis)", "Condicional — só com curvas de nível ou modelo de terreno",
     "Condepe/Fidem — Acervo Cartográfico (curvas de nível); ESIG (a verificar)",
     "Medida mais objetiva de acúmulo de água. 'Relevo' do documento da equipe.",
     "Condepe/Fidem localizado na busca: curvas de 1 m na área urbana (não aberto)."),
    ("Exposição", "E1", "População",
     "Moradores no círculo de 300 m (soma de população × fração da célula dentro do círculo)",
     "Quintil", Q, "Planilha original", "Manter",
     "IBGE — Grade Estatística 2022 (células de 200 × 200 m)",
     "No painel, mostrar também o valor em 1 km.", "Confirmado na busca (não baixado)."),
    ("Exposição", "E2", "Equipamentos públicos",
     "Nº de escolas, creches e unidades de saúde no círculo", "Faixa fixa",
     ("Nenhum", "1", "2", "3–4", "≥ 5 ou hospital/UPA no círculo"),
     "Proposta Claude (provisória)", "Manter",
     "Portal de dados abertos (escolas municipais, UBS, hospitais, rede de saúde); IBGE — CNEFE 2022 (ensino e saúde); CNES/DATASUS",
     "Inclui 'hospitais/escolas atingidos' do documento da equipe (sem contar duas vezes no Impacto). "
     "Perguntar se 'equipamentos públicos' inclui outros (CRAS, abrigos). Calibrar com as contagens reais.",
     "CNEFE confirmado na busca (não baixado). Bases do portal localizadas em 03/10 (não abertas)."),
    ("Exposição", "E3", "Comércio e serviços",
     "Nº de comércios e serviços no círculo", "Quintil", Q, "Planilha original", "Manter",
     "IBGE — CNEFE 2022 (outras finalidades); IPTU (tipo de empreendimento)",
     "'Outras finalidades' pode incluir templos e órgãos públicos. Inclui 'perdas econômicas' do "
     "documento da equipe, que não tem fonte própria.",
     "CNEFE confirmado na busca (não baixado)."),
    ("Vulnerabilidade", "V2", "População vulnerável",
     "% de moradores com 60+ anos ou 0–4 anos nos setores do círculo (ponderado pela área)",
     "Quintil", Q, "Planilha original", "Manter",
     "IBGE — Agregados por setor 2022 (idade)", "",
     "Idade por setor confirmada na busca (não baixado)."),
    ("Vulnerabilidade", "V3", "ZEIS, favelas e comunidades",
     "% da área do círculo em favelas e comunidades urbanas (IBGE 2022) ou em ZEIS 1", "Faixa fixa",
     ("0%", "< 10%", "10–25%", "25–50%", "> 50%"), "Proposta Claude (provisória)",
     "Manter — substitui o V1 se a malha 2022 existir",
     "IBGE — Favelas e Comunidades Urbanas 2022 (malha); ZEIS do Plano Diretor e relação de favelas (portal)",
     "Mais atual que o IVS (2010). ZEIS 2 (terrenos vazios) não conta.",
     "ZEIS e relação de favelas localizadas no portal (não abertas); malha IBGE não verificada."),
    ("Vulnerabilidade", "V4", "Tipo de moradia",
     "% de domicílios do tipo casa nos setores do círculo", "Quintil", Q,
     "Proposta Claude (quintis)", "Manter",
     "IBGE — Agregados por setor 2022 (tipo de domicílio)",
     "Casa térrea é mais atingida que apartamento.",
     "Não verificado se a variável está no arquivo por setor."),
    ("Vulnerabilidade", "V5", "Dificuldade de evacuação",
     "Definição pendente (pendência 15)", "", ("", "", "", "", ""),
     "Documento da equipe (Hellis, 03/10/2026)",
     "Nova — só entra com definição e fonte", "", "Pedir a definição a Hellis antes de procurar dado.",
     "Não pesquisado."),
    ("Vulnerabilidade", "V6", "Infraestrutura precária",
     "Definição pendente (pendência 15)", "", ("", "", "", "", ""),
     "Documento da equipe (Hellis, 03/10/2026)",
     "Nova — só entra com definição e fonte", "",
     "Pode absorver o V4 (tipo de moradia) e o I2 (esgoto). Pedir a definição a Hellis.",
     "Não pesquisado."),
    ("Vulnerabilidade", "V1", "Vulnerabilidade social (IVS)",
     "Maior IVS entre as UDHs que tocam o círculo", "Faixa fixa",
     ("Muito baixa (0–0,200)", "Baixa (0,201–0,300)", "Média (0,301–0,400)",
      "Alta (0,401–0,500)", "Muito alta (> 0,500)"),
     "Faixas do Ipea (conferir no Atlas)", "Plano B — só se o V3 não estiver disponível",
     "Ipea — Atlas da Vulnerabilidade Social (UDH)", "Base do Censo 2010.",
     "Site confirmado na busca; faixas e ano-base a confirmar."),
    ("Impacto", "I1", "Interrupção do trânsito",
     "Classe da via mais importante afetada pelo alagamento; +1 se houver interdição registrada pela CTTU (máximo 5)",
     "Classe",
     ("Via local ou de pedestres", "Coletora", "Arterial secundária", "Arterial principal",
      "Trânsito rápido ou corredor de ônibus"),
     "Proposta Claude, a partir da classificação viária usada na planilha original",
     "Manter",
     "Hierarquia viária (ESIG / CTTU); OSM; corredores de ônibus e fluxo de veículos (portal); paradas (Grande Recife); interdições (CTTU)",
     "Junta 'vias principais' e 'interrupção do trânsito' do documento da equipe. Pedir a Antônio (CTTU): "
     "hierarquia viária, linhas e registro de interdições.",
     "Vias no ESIG confirmadas na busca; ônibus e interdições não verificados."),
    ("Impacto", "I2", "Risco sanitário",
     "% de domicílios sem ligação à rede de esgoto nos setores do círculo", "Quintil", Q,
     "Proposta Claude (quintis)", "Manter",
     "IBGE — Agregados por setor 2022 (esgotamento); Compesa",
     "Quintis porque a cobertura de esgoto no Recife é baixa e faixas fixas podem saturar. "
     "Pode passar para o V6 (infraestrutura precária).",
     "Esgotamento por setor confirmado na busca (não baixado)."),
    ("Impacto", "I3", "Isolamento",
     "Definição pendente (pendência 15)", "", ("", "", "", "", ""),
     "Documento da equipe (Hellis, 03/10/2026)",
     "Nova — só entra com definição e fonte", "",
     "Leitura possível: comunidade sem rota alternativa quando o ponto alaga (malha viária). Confirmar com Hellis.",
     "Não pesquisado."),
]

# (grupo, termo, o que significa, exemplo no projeto)
GLOSSARIO = [
    ("Modelo de risco", "Risco",
     "Medida de quão prioritário é um ponto para a prevenção. Calculado como Probabilidade × Consequência, numa escala de 1 a 25.",
     "Um ponto com Probabilidade 4 e Consequência 3 tem risco 12."),
    ("Modelo de risco", "Critério",
     "Um dos 4 grandes componentes do risco: Probabilidade, Exposição, Vulnerabilidade e Impacto.",
     "Vulnerabilidade é um critério."),
    ("Modelo de risco", "Subcritério",
     "Um item medível dentro de um critério. Cada subcritério recebe uma nota de 1 a 5.",
     "População (E1) é um subcritério da Exposição."),
    ("Modelo de risco", "Probabilidade",
     "Quanto e com que facilidade o ponto alaga: frequência, severidade, maré, proximidade de rio, impermeabilização.",
     "Um ponto que alaga 7 vezes por ano tem Frequência (P1) = 5."),
    ("Modelo de risco", "Exposição",
     "Quem e o que está no caminho da água dentro do círculo de 300 m: moradores, equipamentos públicos, comércio. O que está no círculo conta só aqui, não no Impacto.",
     "Um ponto com muitas escolas por perto tem Equipamentos públicos (E2) alto."),
    ("Modelo de risco", "Vulnerabilidade",
     "Quanto quem está ali tem dificuldade de lidar com o alagamento: idosos e crianças, ZEIS, favelas e comunidades, moradia térrea.",
     "Área com muitas casas térreas tem Tipo de moradia (V4) alto."),
    ("Modelo de risco", "Impacto",
     "Consequência que vai além do local: trânsito nas vias principais, saúde pública (esgoto) e isolamento.",
     "Alagamento numa avenida arterial com ônibus tem Interrupção do trânsito (I1) alta."),
    ("Modelo de risco", "Consequência",
     "Média ponderada de Exposição, Vulnerabilidade e Impacto, de 1 a 5. É o que multiplica a Probabilidade.",
     "Exposição 4, Vulnerabilidade 2 e Impacto 3 com pesos iguais dão Consequência 3."),
    ("Modelo de risco", "Nota (1 a 5)",
     "Valor dado a cada subcritério. 5 = mais crítico, 1 = menos crítico.",
     "Nota 5 em Maré (P3) = o ponto alaga conforme a maré."),
    ("Modelo de risco", "Peso",
     "Importância relativa de um subcritério dentro do critério, ou de um critério dentro da Consequência. Definido pelo grupo. A Probabilidade não tem peso, porque multiplica.",
     "Frequência e Severidade terão peso maior que Proximidade de rio."),
    ("Modelo de risco", "Média ponderada",
     "Média em que cada item conta conforme o seu peso.",
     "Notas 5 (peso 2) e 2 (peso 1): (5×2 + 2×1) ÷ 3 = 4."),
    ("Modelo de risco", "Cobertura",
     "Quantos subcritérios têm nota para o ponto. Nota faltante não vira zero nem média: fica de fora e a cobertura mostra a falta.",
     "Cobertura 11/13 = faltam 2 notas para aquele ponto."),
    ("Modelo de risco", "Sem Probabilidade, sem risco",
     "Como o risco é uma multiplicação, um ponto sem nota de Probabilidade não tem risco calculado nem posição no ranking.",
     "Por isso a ficha com a EMLURB é prioritária."),
    ("Modelo de risco", "Contexto no painel",
     "Informação mostrada no painel sem entrar na nota, porque varia pouco entre os 21 pontos ou não tem dado por ponto.",
     "Intensidade das chuvas (documento da equipe)."),
    ("Modelo de risco", "Tratabilidade",
     "Se é fácil agir no ponto (governança, natureza da causa, obra ou orçamento em andamento). Fica FORA do risco. 5 = mais fácil.",
     "Ponto com dono definido e causa de manutenção tem tratabilidade alta."),
    ("Modelo de risco", "Matriz risco × tratabilidade",
     "Cruza o risco (corte em 9) com a tratabilidade (corte em 3) e coloca cada ponto num de 4 quadrantes.",
     "Risco 12 e tratabilidade 4 = quadrante Agir já."),
    ("Modelo de risco", "Agir já / Estruturar / Oportunidade / Monitorar",
     "Quadrantes da matriz. Agir já: risco alto, fácil de agir. Estruturar: risco alto, difícil (exige articulação). Oportunidade: risco baixo, fácil. Monitorar: risco baixo, difícil.",
     ""),
    ("Modelo de risco", "Diagnóstico de causa",
     "Por que o ponto alaga (drenagem obstruída, subdimensionada, maré, esgoto). Fica fora do risco e indica qual órgão age.",
     "Galeria assoreada → EMLURB (manutenção)."),
    ("Modelo de risco", "Análise de sensibilidade",
     "Recalcular o ranking com outros pesos (ex.: pesos iguais) para ver se o resultado depende demais dos pesos escolhidos.",
     "Se o top 5 muda muito, avisar a alta gestão."),

    ("Tipo de nota", "Faixa fixa",
     "O valor medido cai numa faixa definida antes da coleta. A nota não depende dos outros pontos.",
     "P1: 5 alagamentos/ano cai em '5–6' = nota 4."),
    ("Tipo de nota", "Classe",
     "O ponto é encaixado numa categoria descritiva, não num número.",
     "P3: 'agrava ocasionalmente' = nota 3."),
    ("Tipo de nota", "Quintil",
     "Nota relativa entre os 21 pontos: os valores são ordenados e divididos em 5 grupos de ~4 pontos. Os 20% menores recebem 1; os 20% maiores, 5.",
     "E1: os ~4 pontos com mais moradores recebem nota 5."),
    ("Tipo de nota", "Quintil invertido",
     "Igual ao quintil, mas o MENOR valor recebe a nota mais alta.",
     "P4: o ponto mais perto do rio recebe nota 5."),
    ("Tipo de nota", "Faixa provisória",
     "Faixa proposta antes de ver os dados reais. Deve ser calibrada quando os dados chegarem.",
     "Se todos os pontos caírem na mesma faixa, ela precisa ser ajustada."),
    ("Tipo de nota", "Origem da faixa",
     "De onde veio a escala: planilha original da equipe, referência oficial (ex.: Ipea) ou proposta do Claude (a revisar pelo grupo).",
     ""),

    ("Território e SIG", "Ponto crítico",
     "Local da lista da EMLURB onde há alagamento recorrente. Pode ser um ponto, um trecho de rua ou vários trechos.",
     "Ponto 24 — Ipsep, Rua Blumenau (até a maré)."),
    ("Território e SIG", "Trecho (Linha / Vários)",
     "Ponto crítico que é uma extensão de rua, não um local único. 12 dos 21 pontos são trechos.",
     "Ponto 11 — Rua Professor José Brandão (toda a extensão)."),
    ("Território e SIG", "RPA 6",
     "Região Político-Administrativa 6 do Recife (Boa Viagem, Pina, Imbiribeira, Ipsep, Ibura, Jordão, entre outros bairros). Escopo do piloto.",
     ""),
    ("Território e SIG", "SIG",
     "Sistema de Informação Geográfica: programa para mapas e cálculos espaciais. Usamos o ArcGIS Pro.",
     ""),
    ("Território e SIG", "Círculo de 300 m (buffer)",
     "Área de 300 m em volta do ponto (ou da linha do trecho) onde se medem os subcritérios do SIG. Usado no cálculo do risco.",
     "Moradores dentro do círculo = População (E1)."),
    ("Território e SIG", "Área de influência de 1 km",
     "Círculo de 1 km mostrado no painel só como informação. Não entra no cálculo, porque com 1 km quase todos os pontos se sobrepõem.",
     ""),
    ("Território e SIG", "Grupo de sobreposição (G1–G5)",
     "Pontos cujos círculos de 300 m se sobrepõem. Nesses grupos os subcritérios do SIG tendem a dar notas parecidas.",
     "G1 = pontos 5, 6, 7, 11 e 12 (Conselheiro Aguiar)."),
    ("Território e SIG", "Ponderação por área",
     "Quando uma área de dados (setor, célula) corta o círculo, conta só a parte que fica dentro dele.",
     "Setor com 40% da área no círculo contribui com 40% da população."),
    ("Território e SIG", "EPSG:31985 (SIRGAS 2000 / UTM 25S)",
     "Sistema de coordenadas usado no projeto. Todas as camadas precisam estar nele antes dos cálculos.",
     ""),
    ("Território e SIG", "Coordenada aproximada / verificada",
     "Aproximada = localizada automaticamente e ainda não conferida. Verificada = conferida com a EMLURB ou por imagem de satélite.",
     "Hoje nenhuma das 21 coordenadas está verificada."),
    ("Território e SIG", "Bacia / sub-bacia",
     "Área que drena para um mesmo rio ou canal. Hierarquia do Plano de Drenagem: Bacia → Sub-bacia → Ponto crítico → Obra.",
     "Usada para ligar obras aos pontos."),
    ("Território e SIG", "Shapefile / GeoJSON",
     "Formatos de arquivo de mapa que o ArcGIS abre.",
     "dados/processados/pontos_rpa6.geojson tem os 21 pontos."),

    ("Fontes de dados", "Censo 2022 (IBGE)",
     "Recenseamento de 2022. Fonte de população, idade, tipo de domicílio e esgotamento.", ""),
    ("Fontes de dados", "Setor censitário",
     "Menor área de divulgação do Censo (no Recife adensado, poucos quarteirões).", ""),
    ("Fontes de dados", "Grade Estatística (IBGE)",
     "Malha de quadrados de 200 × 200 m (área urbana) com a população do Censo 2022.", "Fonte do E1."),
    ("Fontes de dados", "CNEFE (IBGE)",
     "Cadastro de endereços do Censo 2022, com coordenada e tipo de uso (domicílio, ensino, saúde, outras finalidades).",
     "Fonte do E2 e do E3."),
    ("Fontes de dados", "Favelas e Comunidades Urbanas (IBGE)",
     "Mapeamento do IBGE das favelas e comunidades urbanas no Censo 2022. Publicação da malha ainda não confirmada.",
     "Fonte do V3."),
    ("Fontes de dados", "IVS (Ipea)",
     "Índice de Vulnerabilidade Social do Ipea, de 0 a 1 (quanto maior, mais vulnerável). Base do Censo 2010.",
     "Plano B do V3."),
    ("Fontes de dados", "UDH",
     "Unidade de Desenvolvimento Humano: área usada pelo Ipea, maior que o setor censitário.", ""),
    ("Fontes de dados", "MapBiomas",
     "Mapa de cobertura do solo (vegetação, área construída etc.) com pixels de 30 m. Não verificado.", "Fonte do P5."),
    ("Fontes de dados", "ESIG",
     "Sistema de informações geográficas da Prefeitura do Recife (vias, lotes, edificações, recursos hídricos).", ""),
    ("Fontes de dados", "CNES / DATASUS",
     "Cadastro Nacional de Estabelecimentos de Saúde.", "Complementa o E2."),
    ("Fontes de dados", "APAC / Cemaden",
     "Agência Pernambucana de Águas e Clima / Centro Nacional de Monitoramento de Desastres: dados de chuva e marés.", ""),
    ("Fontes de dados", "PMDR",
     "Plano Municipal de Drenagem do Recife. Define bacias e sub-bacias.", ""),
    ("Fontes de dados", "Ficha EMLURB",
     "Formulário a ser preenchido com a EMLURB, com uma linha por ponto: localização, frequência, severidade, maré e obras ligadas.",
     "Fonte provisória de P1, P2 e P3."),

    ("Checklist", "Recomendação",
     "Manter = entra no modelo. Condicional = só entra se o dado existir. Plano B = usado só se outro subcritério falhar. Peso alto / menor = sugestão para a definição dos pesos.",
     ""),
    ("Checklist", "Fonte confirmada?",
     "A base de dados existe e a equipe consegue acessá-la (download ou pedido atendido).", ""),
    ("Checklist", "Dado existe na fonte?",
     "A variável necessária está na base, na escala do círculo de 300 m, e cobre os 21 pontos.",
     "O IBGE existe (fonte confirmada), mas falta checar se o tipo de domicílio vem por setor."),
    ("Checklist", "Sim / Parcial / Não / Pendente",
     "Sim = confirmado. Parcial = existe, mas com limitação (escala, cobertura, data). Não = não existe ou não é acessível. Pendente = ainda não checado.",
     ""),
    ("Checklist", "Busca prévia (Claude)",
     "O que foi encontrado em busca na internet. 'Confirmado na busca' = a base aparece como publicada, mas o arquivo não foi aberto.",
     ""),

    ("Órgãos", "COP", "Centro de Operações do Recife: monitoramento e resposta aos eventos.", ""),
    ("Órgãos", "EMLURB", "Empresa de Manutenção e Limpeza Urbana: manutenção da drenagem e lista de pontos críticos.", ""),
    ("Órgãos", "Defesa Civil (Sedec)", "Áreas de risco, alertas e vistorias.", ""),
    ("Órgãos", "URB", "Autarquia de Urbanização do Recife: obras urbanas e de drenagem.", ""),
    ("Órgãos", "CTTU", "Autarquia de Trânsito e Transporte Urbano: trânsito, vias e interdições.", ""),
    ("Órgãos", "Compesa", "Companhia Pernambucana de Saneamento: água e esgoto.", ""),
    ("Órgãos", "ProMorar", "Programa de urbanização de comunidades.", ""),
    ("Órgãos", "SEPLAG / SEPLAN", "Planejamento e gestão. A SEPLAN não é a SEPLAG; o papel dela no projeto ainda será definido.", ""),
    ("Órgãos", "Grande Recife Consórcio", "Consórcio de transporte metropolitano: linhas e paradas de ônibus.", ""),
]

CABECALHO = ["Critério", "Código", "Subcritério", "Como medir", "Tipo de nota",
             "Nota 1", "Nota 2", "Nota 3", "Nota 4", "Nota 5", "Origem da faixa",
             "Recomendação", "Fonte", "Fonte confirmada?", "Dado existe na fonte?",
             "Observação", "Busca prévia (Claude)"]
LARGURAS = (14, 7, 22, 34, 12, 15, 15, 15, 15, 17, 22, 24, 34, 12, 12, 34, 30)
COR_CRITERIO = {"Probabilidade": "DDEBF7", "Exposição": "E2EFDA",
                "Vulnerabilidade": "FCE4D6", "Impacto": "EDE2F6"}
COR_NOTA = ("E2EFDA", "F4F9EE", "FFF9E5", "FDE9D9", "F8CBAD")
COLS_PREENCHER = (14, 15, 16)  # N, O, P


def gerar_glossario(wb):
    ws = wb.create_sheet("Glossário")
    ws["A1"] = "Glossário — termos usados nesta planilha e no modelo de risco"
    ws["A1"].font = Font(name=FONTE, bold=True, size=13)
    ws["A2"] = ("Para quem não participou das discussões do modelo. Os papéis dos órgãos são uma leitura "
                "preliminar, ainda não confirmada com cada órgão.")
    ws["A2"].font = Font(name=FONTE, italic=True, size=9)
    ws.merge_cells("A2:D2")
    for c, titulo in enumerate(("Grupo", "Termo", "O que significa", "Exemplo no projeto"), 1):
        cel = ws.cell(row=4, column=c, value=titulo)
        cel.font = Font(name=FONTE, bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="1F4E78")
        cel.alignment = Alignment(vertical="center", horizontal="center")
    fino = Side(style="thin", color="BFBFBF")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    cores = ("DDEBF7", "E2EFDA", "FCE4D6", "EDE2F6", "FFF2CC", "EDEDED")
    grupos = list(dict.fromkeys(g for g, *_ in GLOSSARIO))
    for i, (grupo, termo, significado, exemplo) in enumerate(GLOSSARIO):
        r = 5 + i
        for c, v in enumerate((grupo, termo, significado, exemplo), 1):
            cel = ws.cell(row=r, column=c, value=v)
            cel.font = Font(name=FONTE, size=10, bold=(c == 2))
            cel.alignment = Alignment(wrap_text=True, vertical="top")
            cel.border = borda
            if c == 1:
                cel.fill = PatternFill("solid", fgColor=cores[grupos.index(grupo) % len(cores)])
    for col, largura in zip("ABCD", (18, 30, 70, 45)):
        ws.column_dimensions[col].width = largura
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:D{4 + len(GLOSSARIO)}"


# Blocos da aba de notas: (critério, [códigos dos subcritérios], tem peso de critério?)
BLOCOS = [
    ("Probabilidade", ["P1", "P2", "P3", "P4", "P5", "P6"], False),
    ("Exposição", ["E1", "E2", "E3"], True),
    ("Vulnerabilidade", ["V2", "V3", "V4", "V5", "V6", "V1"], True),
    ("Impacto", ["I1", "I2", "I3"], True),
]


def gerar_notas(wb):
    """Aba com uma linha por ponto, notas de 1 a 5 e o cálculo do risco."""
    from openpyxl.utils import get_column_letter as L

    ws = wb.create_sheet("Notas dos pontos")
    nomes = {cod: sub for _, cod, sub, *_ in SUBCRITERIOS}
    with open(PONTOS, encoding="utf-8") as f:
        pontos = list(csv.DictReader(f))

    ws["A1"] = "Recife Resiliente — Notas dos pontos críticos da RPA 6 e cálculo do risco"
    ws["A1"].font = Font(name=FONTE, bold=True, size=13)
    ws["A2"] = (
        "Preencha as notas (1 a 5, 5 = mais crítico) nas células amarelas, conforme as escalas da aba "
        "'Checklist subcritérios'. Linha 'Peso': vazia = todos os pesos iguais. Nota em branco fica fora "
        "da média (não vira zero). Sem nota de Probabilidade não há risco. P6 e V1 são condicionais e I3, V5 e "
        "V6 ainda não têm definição: deixe em branco se não forem usados. Colunas cinza são calculadas: não edite. 'Grupo de "
        "sobreposição' (G1–G5) = pontos cujos círculos de 300 m se sobrepõem; vazio = ponto isolado. "
        "Coordenadas em graus decimais (WGS 84), ainda NÃO verificadas: 'Aproximado' e 'Interseção' "
        "foram localizadas automaticamente pelo OpenStreetMap."
    )
    ws["A2"].font = Font(name=FONTE, italic=True, size=9)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 40

    LG, LP, LH = 4, 5, 6  # linhas: grupo, peso, cabeçalho
    r0 = LH + 1
    r1 = r0 + len(pontos) - 1
    branco = Font(name=FONTE, bold=True, color="FFFFFF")
    azul_esc = PatternFill("solid", fgColor="1F4E78")
    cinza = PatternFill("solid", fgColor="EDEDED")
    amarelo = PatternFill("solid", fgColor="FFF2CC")
    fino = Side(style="thin", color="BFBFBF")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def cab(col, texto, fill=azul_esc, font=branco, linha=LH):
        c = ws.cell(row=linha, column=col, value=texto)
        c.fill, c.font, c.alignment, c.border = fill, font, centro, borda
        return c

    ident = [("ID", "ponto_id", 6), ("Bairro", "bairro", 13), ("Trecho", "trecho", 42),
             ("Tipo", "tipo", 8), ("Grupo de\nsobreposição", "grupo_sobreposicao", 12),
             ("Latitude", "lat", 11), ("Longitude", "lon", 11),
             ("Precisão da\ncoordenada", "precisao_coordenada", 12)]
    for j, (titulo, _, larg) in enumerate(ident, 1):
        cab(j, titulo)
        ws.column_dimensions[L(j)].width = larg
    ws.cell(row=LP, column=len(ident), value="Peso →").font = Font(name=FONTE, bold=True, size=9)
    ws.cell(row=LP, column=len(ident)).alignment = Alignment(horizontal="right")

    col = len(ident) + 1
    nota_criterio = {}   # critério -> coluna da nota do critério
    notas_cols = []      # colunas de notas de subcritério
    for crit, codigos, tem_peso in BLOCOS:
        cor = PatternFill("solid", fgColor=COR_CRITERIO[crit])
        ini = col
        for cod in codigos:
            cab(col, f"{cod}\n{nomes[cod]}", fill=cor, font=Font(name=FONTE, bold=True, size=9))
            p = ws.cell(row=LP, column=col)
            p.fill, p.border, p.alignment = amarelo, borda, centro
            p.font = Font(name=FONTE, size=9, color="0000FF")
            ws.column_dimensions[L(col)].width = 11
            notas_cols.append(col)
            col += 1
        fim = col - 1
        cab(col, f"Nota\n{crit}", fill=PatternFill("solid", fgColor="595959"))
        ws.column_dimensions[L(col)].width = 13
        if tem_peso:
            p = ws.cell(row=LP, column=col)
            p.fill, p.border, p.alignment = amarelo, borda, centro
            p.font = Font(name=FONTE, size=9, color="0000FF", bold=True)
        else:
            ws.cell(row=LP, column=col, value="(multiplica)").font = Font(name=FONTE, size=8, italic=True)
            ws.cell(row=LP, column=col).alignment = centro
        g = ws.cell(row=LG, column=ini, value=crit)
        g.font, g.fill, g.alignment = Font(name=FONTE, bold=True), cor, centro
        ws.merge_cells(start_row=LG, start_column=ini, end_row=LG, end_column=col)
        for r in range(r0, r1 + 1):
            a, b = f"{L(ini)}{r}", f"{L(fim)}{r}"
            pa, pb = f"{L(ini)}${LP}", f"{L(fim)}${LP}"
            peso = f"({pa}:{pb}+({pa}:{pb}=\"\"))"
            ws[f"{L(col)}{r}"] = (
                f'=IF(COUNT({a}:{b})=0,"",IFERROR(SUMPRODUCT({a}:{b},{peso})'
                f'/SUMPRODUCT(({a}:{b}<>"")*{peso}),""))'
            )
        nota_criterio[crit] = col
        col += 1

    # Resultado
    cC, cR, cI, cPos, cCob = col, col + 1, col + 2, col + 3, col + 4
    resultado = [(cC, "Consequência\n(1–5)"), (cR, "Risco\nP × C (1–25)"),
                 (cI, "Índice de risco\n(0–1)"), (cPos, "Posição"), (cCob, "Notas\npreenchidas")]
    for c, t in resultado:
        cab(c, t, fill=PatternFill("solid", fgColor="C00000"))
        ws.column_dimensions[L(c)].width = 13
    g = ws.cell(row=LG, column=cC, value="Resultado")
    g.font, g.fill, g.alignment = Font(name=FONTE, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="C00000"), centro
    ws.merge_cells(start_row=LG, start_column=cC, end_row=LG, end_column=cCob)

    cons = [nota_criterio[c] for c in ("Exposição", "Vulnerabilidade", "Impacto")]
    cp = L(nota_criterio["Probabilidade"])
    for r in range(r0, r1 + 1):
        w = {c: f"({L(c)}${LP}+({L(c)}${LP}=\"\"))" for c in cons}
        num = "+".join(f"N({L(c)}{r})*{w[c]}" for c in cons)
        den = "+".join(f"ISNUMBER({L(c)}{r})*{w[c]}" for c in cons)
        ws[f"{L(cC)}{r}"] = f'=IFERROR(IF(({den})=0,"",({num})/({den})),"")'
        ws[f"{L(cR)}{r}"] = f'=IF(AND(ISNUMBER({cp}{r}),ISNUMBER({L(cC)}{r})),{cp}{r}*{L(cC)}{r},"")'
        ws[f"{L(cI)}{r}"] = f'=IF(ISNUMBER({L(cR)}{r}),({L(cR)}{r}-1)/24,"")'
        ws[f"{L(cPos)}{r}"] = f'=IF(ISNUMBER({L(cR)}{r}),COUNTIF({L(cR)}${r0}:{L(cR)}${r1},">"&{L(cR)}{r})+1,"")'
        cols_n = ",".join(f"{L(c)}{r}" for c in notas_cols)
        ws[f"{L(cCob)}{r}"] = f"=COUNT({cols_n})"

    # Corpo da tabela
    for i, pt in enumerate(pontos):
        r = r0 + i
        for j, (_, chave, _) in enumerate(ident, 1):
            v = pt[chave]
            if chave == "ponto_id":
                v = int(v)
            elif chave in ("lat", "lon"):
                v = float(v)
            c = ws.cell(row=r, column=j, value=v)
            if chave in ("lat", "lon"):
                c.number_format = "0.000000"
            c.font = Font(name=FONTE, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=(chave == "trecho"))
            c.border = borda
        for c in notas_cols:
            cel = ws.cell(row=r, column=c)
            cel.fill, cel.border, cel.alignment = amarelo, borda, centro
            cel.font = Font(name=FONTE, size=10, color="0000FF")
        for c in list(nota_criterio.values()) + [cC, cR, cI, cPos, cCob]:
            cel = ws.cell(row=r, column=c)
            cel.fill, cel.border, cel.alignment = cinza, borda, centro
            cel.font = Font(name=FONTE, size=10, bold=c in (cR, cI, cPos))
            cel.number_format = "0.00" if c not in (cPos, cCob) else "0"
        ws.cell(row=r, column=cI).number_format = "0.000"

    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                        showErrorMessage=True, errorTitle="Nota inválida",
                        error="Use um número inteiro de 1 a 5 (5 = mais crítico).")
    ws.add_data_validation(dv)
    for c in notas_cols:
        dv.add(f"{L(c)}{r0}:{L(c)}{r1}")
    dvp = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                         showErrorMessage=True, error="Peso deve ser um número maior ou igual a zero.")
    ws.add_data_validation(dvp)
    for c in notas_cols + cons:
        dvp.add(f"{L(c)}{LP}")
    ws.conditional_formatting.add(
        f"{L(cI)}{r0}:{L(cI)}{r1}",
        ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                       mid_type="num", mid_value=8 / 24, mid_color="FFEB84",
                       end_type="num", end_value=1, end_color="F8696B"))

    nota = r1 + 2
    textos = [
        "Como o cálculo funciona:",
        "• Nota do critério = média ponderada das notas dos subcritérios preenchidos (pesos da linha 5; vazio = 1).",
        "• Consequência = média ponderada das notas de Exposição, Vulnerabilidade e Impacto (pesos da linha 5 nas colunas cinza 'Nota').",
        "• Risco = nota da Probabilidade × Consequência (1 a 25).",
        "• Índice de risco (0–1) = (Risco − 1) ÷ 24. 0 = risco mínimo (P = 1 e C = 1); 1 = risco máximo (P = 5 e C = 5). "
        "Corte da matriz (risco 9) = índice 0,333. O ranking é o mesmo do Risco.",
        "• Posição = ranking pelo Risco (1 = maior risco). Empates recebem a mesma posição.",
    ]
    for k, t in enumerate(textos):
        c = ws.cell(row=nota + k, column=1, value=t)
        c.font = Font(name=FONTE, size=9, bold=(k == 0))

    ws.freeze_panes = ws.cell(row=r0, column=len(ident) + 1)
    ws.row_dimensions[LH].height = 54
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return r0, r1, cR, cI


def gerar(saida=SAIDA):
    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist subcritérios"

    ws["A1"] = "Recife Resiliente — Subcritérios, notas, fontes e checklist dos dados (RPA 6)"
    ws["A1"].font = Font(name=FONTE, bold=True, size=13)
    ws["A2"] = (
        "Proposta de 02/10/2026 (Pedro + Claude), atualizada em 03/10/2026 com o documento da equipe "
        "(Hellis); não validada pelo grupo. I3, V5 e V6 são novas e ainda sem definição nem escala. Nota 5 = mais crítico. "
        "Faixa fixa = valor absoluto (provisória: calibrar com os dados reais). Quintil = nota relativa "
        "entre os 21 pontos (Quintil 1 = 20% com menor valor). Preencha as colunas amarelas (N, O e P): "
        "Sim / Não / Parcial / Pendente. 'Fonte confirmada' = a base existe e é acessível; 'Dado existe "
        "na fonte' = a variável está lá, na escala do círculo de 300 m, e cobre os 21 pontos. "
        "Termos explicados na aba Glossário."
    )
    ws["A2"].font = Font(name=FONTE, italic=True, size=9)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:Q2")
    ws.row_dimensions[2].height = 44

    h = 4
    for c, titulo in enumerate(CABECALHO, 1):
        cel = ws.cell(row=h, column=c, value=titulo)
        cel.font = Font(name=FONTE, bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="1F4E78")
        cel.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    fino = Side(style="thin", color="BFBFBF")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    amarelo = PatternFill("solid", fgColor="FFF2CC")
    r0 = h + 1
    for i, (crit, cod, sub, medir, tipo, notas, origem, rec, fonte, obs, busca) in enumerate(SUBCRITERIOS):
        r = r0 + i
        valores = [crit, cod, sub, medir, tipo, *notas, origem, rec, fonte,
                   "Pendente", "Pendente", obs, busca]
        for c, v in enumerate(valores, 1):
            cel = ws.cell(row=r, column=c, value=v)
            cel.border = borda
            cel.font = Font(name=FONTE, size=10, color="0000FF" if c in COLS_PREENCHER else "000000")
            centro = c in (2, 5, 14, 15) or 6 <= c <= 10
            cel.alignment = Alignment(wrap_text=True, vertical="top",
                                      horizontal="center" if centro else "left")
            if c in COLS_PREENCHER:
                cel.fill = amarelo
            elif c in (1, 2):
                cel.fill = PatternFill("solid", fgColor=COR_CRITERIO[crit])
            elif 6 <= c <= 10:
                cel.fill = PatternFill("solid", fgColor=COR_NOTA[c - 6])
    ultima = r0 + len(SUBCRITERIOS) - 1

    dv = DataValidation(type="list", formula1='"Sim,Não,Parcial,Pendente"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"N{r0}:O{ultima}")
    for valor, cor in (("Sim", "C6EFCE"), ("Não", "FFC7CE"), ("Parcial", "FFEB9C")):
        ws.conditional_formatting.add(
            f"N{r0}:O{ultima}",
            CellIsRule(operator="equal", formula=[f'"{valor}"'],
                       fill=PatternFill("solid", fgColor=cor)))

    s = ultima + 2
    ws.cell(row=s, column=1, value="Resumo").font = Font(name=FONTE, bold=True)
    ws.cell(row=s, column=14, value="Fonte confirmada").font = Font(name=FONTE, bold=True, size=9)
    ws.cell(row=s, column=15, value="Dado existe").font = Font(name=FONTE, bold=True, size=9)
    for k, rotulo in enumerate(("Sim", "Parcial", "Não", "Pendente")):
        rr = s + 1 + k
        cel = ws.cell(row=rr, column=13, value=rotulo)
        cel.font = Font(name=FONTE, size=10)
        cel.alignment = Alignment(horizontal="right")
        for col in ("N", "O"):
            ws[f"{col}{rr}"] = f"=COUNTIF({col}${r0}:{col}${ultima},$M{rr})"
            ws[f"{col}{rr}"].font = Font(name=FONTE, size=10)
            ws[f"{col}{rr}"].alignment = Alignment(horizontal="center")
    rr = s + 5
    cel = ws.cell(row=rr, column=13, value="Subcritérios com fonte e dado confirmados")
    cel.font = Font(name=FONTE, bold=True, size=10)
    cel.alignment = Alignment(horizontal="right", wrap_text=True)
    ws[f"O{rr}"] = (f'=COUNTIFS(N{r0}:N{ultima},"Sim",O{r0}:O{ultima},"Sim")'
                    f'&" de "&ROWS(N{r0}:N{ultima})')
    ws[f"O{rr}"].font = Font(name=FONTE, bold=True, size=10)
    ws[f"O{rr}"].alignment = Alignment(horizontal="center")

    for i, largura in enumerate(LARGURAS):
        ws.column_dimensions[chr(ord("A") + i)].width = largura
    ws.freeze_panes = f"D{r0}"
    ws.auto_filter.ref = f"A{h}:Q{ultima}"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    gerar_notas(wb)
    gerar_glossario(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(saida)
    return r0, ultima


if __name__ == "__main__":
    print(gerar())
